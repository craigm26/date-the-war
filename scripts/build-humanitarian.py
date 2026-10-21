#!/usr/bin/env python3
"""Build data/humanitarian.json from the latest INFORM Severity Index release.

INFORM Severity is published monthly by ACAPS on HDX under CC BY. It rates each
humanitarian crisis 1 to 5 (Very low to Very high) from three dimensions:
impact of the crisis, conditions of the people affected, and the complexity of
the crisis. This script reads the country-level sheet and keeps, per country:
the index and category, the 3-month trend, ACAPS's data-reliability rating,
the crisis drivers, people in need (levels 3 to 5), people displaced, and the
ACAPS crisis page.

Needs openpyxl and pycountry (pip install openpyxl pycountry); run monthly:
    python3 scripts/build-humanitarian.py            # newest release on HDX
    python3 scripts/build-humanitarian.py file.xlsx  # a downloaded release
"""
import io, json, re, sys, urllib.request, datetime, pathlib
import openpyxl, pycountry

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATASET = 'inform-global-crisis-severity-index'
PAGE = f'https://data.humdata.org/dataset/{DATASET}'
UA = {'user-agent': 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36 date-the-war'}


def fetch(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60) as r:
        return r.read()


def latest_release():
    meta = json.loads(fetch(f'https://data.humdata.org/api/3/action/package_show?id={DATASET}'))['result']
    res = [r for r in meta['resources'] if re.match(r'^\d{6} INFORM Severity', r['name']) and r['format'].upper() == 'XLSX']
    res.sort(key=lambda r: r['name'][:6], reverse=True)
    return res[0]['name'], res[0]['url']


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def header_index(rows, needle, row=None):
    """Column index of the first header cell containing `needle` in the given header row."""
    for i, c in enumerate(rows[row]):
        if isinstance(c, str) and needle.lower() in c.lower():
            return i
    raise KeyError(needle)


def main():
    if len(sys.argv) > 1:
        name, data = pathlib.Path(sys.argv[1]).name, pathlib.Path(sys.argv[1]).read_bytes()
    else:
        name, url = latest_release()
        data = fetch(url)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)

    about = [r[0] for r in wb['About'].iter_rows(values_only=True) if r and r[0]]
    released = next((datetime.datetime.strptime(a, '%d/%m/%Y').date().isoformat() for a in about if isinstance(a, str) and re.match(r'\d\d/\d\d/\d{4}$', a)), None)
    edition = next((a for a in about if isinstance(a, str) and re.match(r'^[A-Z][a-z]+ \d{4}$', a)), name)

    # crisis-level indicators: displaced and people in each condition level
    ind = list(wb['Crisis Indicator Data'].iter_rows(values_only=True))
    hid = 0
    c_id = header_index(ind, 'CRISIS ID', hid)
    c_disp = header_index(ind, 'displaced people', hid)
    heads = [c if isinstance(c, str) else '' for c in ind[hid]]
    level_cols = {}
    for i, h in enumerate(heads):
        m = re.search(r'level (\d)', h.lower())
        if m and h.lower().startswith('# of people'):
            level_cols[int(m.group(1))] = i
    indicators = {}
    for r in ind[2:]:
        if not r[c_id]:
            continue
        pin_parts = [num(r[level_cols[k]]) for k in (3, 4, 5) if k in level_cols]
        indicators[r[c_id]] = {
            'displaced': num(r[c_disp]),
            'pin': sum(p for p in pin_parts if p) if any(p is not None for p in pin_parts) else None,
        }

    info = {}
    for r in list(wb['Crisis info'].iter_rows(values_only=True))[1:]:
        if r[0]:
            info[r[0]] = {'url': (r[10] or '').replace(' ', '%20'), 'description': r[9]}

    rows = [r for r in list(wb['INFORM Severity - country'].iter_rows(values_only=True))[4:] if r[0] and r[1]]
    out = []
    for r in rows:
        crisis, cid, country, iso3, drivers, index, cat, label, trend, rel = r[:10]
        region = r[19]
        c = pycountry.countries.get(alpha_3=iso3)
        out.append({
            'iso3': iso3,
            'isoNum': c.numeric if c else None,
            'country': country,
            'region': region,
            'crisis': crisis.strip(),
            'crisisId': cid,
            'index': round(float(index), 1),
            'category': int(cat),
            'label': label,
            'trend': trend,
            'reliability': rel,
            'drivers': [d.strip() for d in (drivers or '').split(',') if d.strip()],
            'pin': indicators.get(cid, {}).get('pin'),
            'displaced': indicators.get(cid, {}).get('displaced'),
            'url': info.get(cid, {}).get('url') or 'https://www.acaps.org/en/countries',
        })
    out.sort(key=lambda x: (-x['index'], x['country']))
    doc = {
        'source': {
            'name': 'INFORM Severity Index', 'publisher': 'ACAPS', 'edition': edition, 'released': released,
            'url': PAGE, 'license': 'CC BY 4.0', 'file': name,
            'built': datetime.date.today().isoformat(),
        },
        'countries': out,
    }
    path = ROOT / 'data' / 'humanitarian.json'
    path.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + '\n')
    print(f'wrote {len(out)} countries from {name} ({edition}, released {released})')


if __name__ == '__main__':
    main()
